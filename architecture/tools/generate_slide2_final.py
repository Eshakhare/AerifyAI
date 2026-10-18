import pptx
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

pptx_path = r'd:\olderEdrive\EshaProjects\AerifyAI\AerifyAI_AI_Builder_Cup_Pitch_Deck.pptx'
prs = pptx.Presentation(pptx_path)
s2 = prs.slides[2]

# Remove extra shapes added previously beyond base 4 (header banner, bottom line, title, light canvas)
while len(s2.shapes) > 4:
    sp = s2.shapes[-1]
    sp._element.getparent().remove(sp._element)

# Ensure title is positioned cleanly on light canvas
for shape in s2.shapes:
    if shape.has_text_frame and 'Brief about the idea' in shape.text:
        shape.left = Emu(400000)
        shape.top = Emu(560000)
        shape.width = Emu(8344000)
        shape.height = Emu(650000)
        tf = shape.text_frame
        tf.clear()
        p0 = tf.paragraphs[0]
        p0.text = 'Brief about the idea'
        p0.font.name = 'Google Sans'
        p0.font.size = Pt(21)
        p0.font.bold = True
        p0.font.color.rgb = RGBColor(15, 23, 42) # Slate 900
        
        p1 = tf.add_paragraph()
        p1.text = 'EXECUTIVE SUMMARY & NATIONAL PLATFORM SCOPE'
        p1.font.name = 'Google Sans'
        p1.font.size = Pt(8.5)
        p1.font.bold = True
        p1.font.color.rgb = RGBColor(100, 116, 139) # Slate 500

# 3-Card layout dimensions (side-by-side)
card_lefts = [Emu(400000), Emu(3248000), Emu(6096000)]
card_top = Emu(1260000)
card_w = Emu(2648000)
card_h = Emu(3380000)

three_cards = [
    {
        'tag_icon': '⚠️',
        'subtitle': '01 // THE PROBLEM',
        'title': 'Urban Smog & Zero Proof',
        'theme_color': RGBColor(245, 158, 11),   # Google Amber/Yellow (#F59E0B)
        'card_bg': RGBColor(254, 252, 232),       # Light Pastel Yellow (#FEFCE8)
        'header_color': RGBColor(180, 83, 9),     # Deep Amber (#B45309)
        'border_width': Pt(1.75),
        'bullets': [
            ('Station Blind Spots: ', 'Only ~40 towers for 30M+ people; vast neighborhoods unmonitored.'),
            ('Unverified Spending: ', 'Crores spent on sprinklers with zero proof of actual impact.'),
            ('Weather Confusion: ', 'Cannot separate clean-up results from ambient wind or rain shifts.'),
            ('Short-Lived Impact: ', 'No tracking to measure how quickly pollution rebounds after action.'),
            ('Zero Civic Proof: ', 'Citizens see scary alerts with no visibility into local municipal work.')
        ]
    },
    {
        'tag_icon': '🌐',
        'subtitle': '02 // THE SOLUTION (HERO)',
        'title': 'Action-Driven Platform',
        'theme_color': RGBColor(66, 133, 244),   # Google Blue (#4285F4)
        'card_bg': RGBColor(240, 247, 255),       # Light Pastel Blue (#F0F7FF)
        'header_color': RGBColor(26, 115, 232),   # Deep Blue (#1A73E8)
        'border_width': Pt(2.25),                # Accent border for Hero card
        'bullets': [
            ('Continuous 500m Grid: ', 'Google satellite telemetry covers every street & ward continuously.'),
            ('Measures Unmeasured Tasks: ', 'Calculates exact before/after AQI drop for every field action.'),
            ('Causal DiD Engine: ', 'Compares cleaned vs. untouched spots to eliminate weather bias.'),
            ('4-Timed Checkpoints: ', 'Auto-samples air at +15s, +3h, +6h, +24h (configurable to 3 days).'),
            ('Public Verification Ledger: ', 'Publishes transparent neighborhood proof to restore civic trust.')
        ]
    },
    {
        'tag_icon': '👥',
        'subtitle': '03 // TARGET AUDIENCE',
        'title': 'Citizens, Officers & Cities',
        'theme_color': RGBColor(52, 168, 83),    # Google Green (#34A853)
        'card_bg': RGBColor(240, 253, 244),       # Light Pastel Green (#F0FDF4)
        'header_color': RGBColor(21, 128, 61),    # Deep Green (#15803D)
        'border_width': Pt(1.75),
        'bullets': [
            ('General Public (No Login): ', 'Live street AQI, pollen risk, and state directives (Odd-Even, bans).'),
            ('Citizens (Google Login): ', 'Pin & monitor places you care about for 30-day verified action proof.'),
            ('Municipal Officers (Gov): ', 'Dispatch field teams, set 4-interval timers, and audit action efficacy.'),
            ('Regulators & Leadership: ', 'Empirically verify clean-air spending with causal DiD evidence trails.'),
            ('Zero Privacy Tracking: ', 'Strictly local pin-based monitoring with zero background GPS tracking.')
        ]
    }
]

for idx, c in enumerate(three_cards):
    left = card_lefts[idx]
    
    # Background Card shape
    card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, card_top, card_w, card_h)
    card.fill.solid()
    card.fill.fore_color.rgb = c['card_bg']
    card.line.color.rgb = c['theme_color']
    card.line.width = c['border_width']
    
    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = Pt(11)
    tf.margin_top = Pt(11)
    tf.margin_right = Pt(11)
    tf.margin_bottom = Pt(8)
    
    # Subtitle / Category Tag
    p0 = tf.paragraphs[0]
    p0.text = f"{c['tag_icon']}  {c['subtitle']}"
    p0.font.name = 'Google Sans'
    p0.font.size = Pt(8.5)
    p0.font.bold = True
    p0.font.color.rgb = c['header_color']
    p0.space_after = Pt(2)
    
    # Main Headline
    p1 = tf.add_paragraph()
    p1.text = c['title']
    p1.font.name = 'Google Sans'
    p1.font.size = Pt(12.5)
    p1.font.bold = True
    p1.font.color.rgb = RGBColor(15, 23, 42)
    p1.space_after = Pt(6)
    
    # 5 Bullets
    for prefix, body in c['bullets']:
        pb = tf.add_paragraph()
        pb.space_after = Pt(4)
        
        r_dot = pb.add_run()
        r_dot.text = '• '
        r_dot.font.name = 'Google Sans'
        r_dot.font.size = Pt(9)
        r_dot.font.bold = True
        r_dot.font.color.rgb = c['theme_color']
        
        r_pre = pb.add_run()
        r_pre.text = prefix
        r_pre.font.name = 'Google Sans'
        r_pre.font.size = Pt(8.8)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(30, 41, 59)
        
        r_txt = pb.add_run()
        r_txt.text = body
        r_txt.font.name = 'Google Sans'
        r_txt.font.size = Pt(8.8)
        r_txt.font.color.rgb = RGBColor(71, 85, 105)

# Bottom Roadmap Banner in light theme
b_left = Emu(400000)
b_top = Emu(4700000)
b_w = Emu(8344000)
b_h = Emu(320000)

banner = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, b_left, b_top, b_w, b_h)
banner.fill.solid()
banner.fill.fore_color.rgb = RGBColor(255, 255, 255)
banner.line.color.rgb = RGBColor(66, 133, 244) # Google Blue
banner.line.width = Pt(1.25)

btf = banner.text_frame
btf.margin_left = Pt(12)
btf.margin_top = Pt(4)
btf.margin_right = Pt(12)
btf.margin_bottom = Pt(4)

bp = btf.paragraphs[0]
r_tag = bp.add_run()
r_tag.text = '💡 Live National Deployment | '
r_tag.font.name = 'Google Sans'
r_tag.font.size = Pt(9)
r_tag.font.bold = True
r_tag.font.color.rgb = RGBColor(26, 115, 232)

r_desc = bp.add_run()
r_desc.text = 'Auto-Detects State & District via GPS • Multi-State Selector & Reset • 500m Satellite Telemetry Across India'
r_desc.font.name = 'Google Sans'
r_desc.font.size = Pt(9)
r_desc.font.color.rgb = RGBColor(51, 65, 85)

prs.save(pptx_path)
print('Slide 2 exclusively regenerated and saved in PPTX!')
