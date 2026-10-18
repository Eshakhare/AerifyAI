import pptx
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

pptx_path = r'd:\olderEdrive\EshaProjects\AerifyAI\AerifyAI_AI_Builder_Cup_Pitch_Deck.pptx'
prs = pptx.Presentation(pptx_path)

# Slide 8 is at index 8 (Technologies to be used in the solution)
s8 = prs.slides[8]

# Remove shapes beyond base 4 shapes
while len(s8.shapes) > 4:
    sp = s8.shapes[-1]
    sp._element.getparent().remove(sp._element)

# Add full-width light canvas
canvas = s8.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(0), Emu(507359), Emu(9144000), Emu(4636141))
canvas.fill.solid()
canvas.fill.fore_color.rgb = RGBColor(248, 250, 252) # Soft slate-50
canvas.line.fill.background()

# Style Title Text Box (Shape 3)
for shape in s8.shapes:
    if shape.has_text_frame and ('technologies' in shape.text.lower() or 'tech' in shape.text.lower()):
        shape.left = Emu(350000)
        shape.top = Emu(550000)
        shape.width = Emu(8444000)
        shape.height = Emu(620000)
        tf = shape.text_frame
        tf.clear()
        p0 = tf.paragraphs[0]
        p0.text = 'Technologies to be used in the solution'
        p0.font.name = 'Google Sans'
        p0.font.size = Pt(20)
        p0.font.bold = True
        p0.font.color.rgb = RGBColor(15, 23, 42)
        
        p1 = tf.add_paragraph()
        p1.text = 'COMPACT TECHNOLOGY STACK: CATEGORIES, TOOLS & ONE-LINE PURPOSES'
        p1.font.name = 'Google Sans'
        p1.font.size = Pt(8)
        p1.font.bold = True
        p1.font.color.rgb = RGBColor(100, 116, 139)

# ------------------------------------------------------------------------------
# 6 COMPACT TECHNOLOGY CARDS (3 columns x 2 rows)
# ------------------------------------------------------------------------------
col_lefts = [Emu(350000), Emu(3218000), Emu(6086000)]
row_tops = [Emu(1240000), Emu(3040000)]
card_w = Emu(2708000)
card_h = Emu(1680000)

tech_cards = [
    # Row 0, Col 0: Frontend & Mapping
    {
        'col': 0, 'row': 0,
        'tag': '01 // FRONTEND & MAPPING',
        'theme_color': RGBColor(66, 133, 244), # Google Blue
        'card_bg': RGBColor(240, 247, 255),
        'tools': 'React 19 · Vite · Maps JS SDK',
        'purpose': 'Renders interactive 3D vector map, persistent Near You street cards, and SVG rolling trend line graphs.',
        'tech_chips': 'TypeScript · Tailwind CSS · Lucide'
    },
    # Row 0, Col 1: AI / GenAI
    {
        'col': 1, 'row': 0,
        'tag': '02 // AI / GENAI LAYER',
        'theme_color': RGBColor(124, 58, 237), # Purple
        'card_bg': RGBColor(245, 243, 255),
        'tools': 'Vertex AI · Gemini 2.5 Flash-Lite',
        'purpose': 'Generates plain-English executive briefings, plume explanations, and verifies statutory gazette circulars.',
        'tech_chips': 'google-genai SDK · Structured JSON · MCP'
    },
    # Row 0, Col 2: Environmental Data
    {
        'col': 2, 'row': 0,
        'tag': '03 // ENVIRONMENTAL TELEMETRY',
        'theme_color': RGBColor(52, 168, 83), # Google Green
        'card_bg': RGBColor(240, 253, 244),
        'tools': 'Air Quality · Pollen · Weather APIs',
        'purpose': 'Provides continuous 500m ground truth, 30-day hourly history, and 3-day multi-allergen pollen forecasts.',
        'tech_chips': 'Copernicus Sentinel-5P · Geocoding API'
    },
    # Row 1, Col 0: Backend & Compute
    {
        'col': 0, 'row': 1,
        'tag': '04 // BACKEND & COMPUTE',
        'theme_color': RGBColor(217, 119, 6), # Google Amber
        'card_bg': RGBColor(254, 252, 232),
        'tools': 'Python 3.12 · FastAPI · Cloud Run',
        'purpose': 'Powers asynchronous microservices, Difference-in-Differences causal workers, and 4-stage timers.',
        'tech_chips': 'SQLAlchemy 2.0 · Pydantic · Docker'
    },
    # Row 1, Col 1: Data & Spatial Engine
    {
        'col': 1, 'row': 1,
        'tag': '05 // DATA & SPATIAL ENGINE',
        'theme_color': RGBColor(234, 67, 53), # Google Red
        'card_bg': RGBColor(254, 242, 242),
        'tools': 'Cloud SQL · PostgreSQL 16 · PostGIS',
        'purpose': 'Stores authoritative municipal action plans, checkpoint states, and geographic district GeoJSON polygons.',
        'tech_chips': 'PostGIS Spatial · Alembic Migrations'
    },
    # Row 1, Col 2: Security & Authentication
    {
        'col': 2, 'row': 1,
        'tag': '06 // SECURITY & AUTHENTICATION',
        'theme_color': RGBColor(14, 116, 144), # Cyan / Navy
        'card_bg': RGBColor(240, 249, 255),
        'tools': 'Firebase Auth · Cloud IAM (Zero-Pwd)',
        'purpose': 'Enforces dual-path RBAC: Google login for citizen pinning and official department email login for officers.',
        'tech_chips': 'google-cloud-sql-connector · RBAC'
    }
]

for c in tech_cards:
    left = col_lefts[c['col']]
    top = row_tops[c['row']]
    
    card = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, card_w, card_h)
    card.fill.solid()
    card.fill.fore_color.rgb = c['card_bg']
    card.line.color.rgb = c['theme_color']
    card.line.width = Pt(1.75)
    
    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = Pt(10)
    tf.margin_top = Pt(8)
    tf.margin_right = Pt(10)
    tf.margin_bottom = Pt(6)
    
    # Tag
    p_tag = tf.paragraphs[0]
    p_tag.text = c['tag']
    p_tag.font.name = 'Google Sans'
    p_tag.font.size = Pt(7.5)
    p_tag.font.bold = True
    p_tag.font.color.rgb = c['theme_color']
    p_tag.space_after = Pt(2)
    
    # Main Tools Headline
    p_tools = tf.add_paragraph()
    p_tools.text = c['tools']
    p_tools.font.name = 'Google Sans'
    p_tools.font.size = Pt(11)
    p_tools.font.bold = True
    p_tools.font.color.rgb = RGBColor(15, 23, 42)
    p_tools.space_after = Pt(4)
    
    # One-Line Purpose
    p_purp = tf.add_paragraph()
    r_purp_lbl = p_purp.add_run()
    r_purp_lbl.text = "Purpose: "
    r_purp_lbl.font.name = 'Google Sans'
    r_purp_lbl.font.size = Pt(8.5)
    r_purp_lbl.font.bold = True
    r_purp_lbl.font.color.rgb = RGBColor(30, 41, 59)
    
    r_purp_txt = p_purp.add_run()
    r_purp_txt.text = c['purpose']
    r_purp_txt.font.name = 'Google Sans'
    r_purp_txt.font.size = Pt(8.5)
    r_purp_txt.font.color.rgb = RGBColor(71, 85, 105)
    p_purp.space_after = Pt(4)
    
    # Chips / Key libraries
    p_chip = tf.add_paragraph()
    r_chip_lbl = p_chip.add_run()
    r_chip_lbl.text = "Key Libs: "
    r_chip_lbl.font.name = 'Google Sans'
    r_chip_lbl.font.size = Pt(7.5)
    r_chip_lbl.font.bold = True
    r_chip_lbl.font.color.rgb = c['theme_color']
    
    r_chip_txt = p_chip.add_run()
    r_chip_txt.text = c['tech_chips']
    r_chip_txt.font.name = 'Google Sans'
    r_chip_txt.font.size = Pt(7.5)
    r_chip_txt.font.color.rgb = RGBColor(100, 116, 139)

# Bottom Banner: Google Cloud & Maps Native Stack Summary
b_left = Emu(350000)
b_top = Emu(4800000)
b_w = Emu(8444000)
b_h = Emu(250000)

banner = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, b_left, b_top, b_w, b_h)
banner.fill.solid()
banner.fill.fore_color.rgb = RGBColor(255, 255, 255)
banner.line.color.rgb = RGBColor(66, 133, 244)
banner.line.width = Pt(1.2)

btf = banner.text_frame
btf.margin_left = Pt(10)
btf.margin_top = Pt(3)
btf.margin_right = Pt(10)
btf.margin_bottom = Pt(3)

bp = btf.paragraphs[0]
r_tag = bp.add_run()
r_tag.text = '☁️ Google Cloud Native | '
r_tag.font.name = 'Google Sans'
r_tag.font.size = Pt(8.5)
r_tag.font.bold = True
r_tag.font.color.rgb = RGBColor(26, 115, 232)

r_desc = bp.add_run()
r_desc.text = 'Zero Static Database Passwords (IAM Connector) • Serverless Cloud Run • 100% Contract Validated (validate_contracts.py)'
r_desc.font.name = 'Google Sans'
r_desc.font.size = Pt(8.5)
r_desc.font.color.rgb = RGBColor(51, 65, 85)

prs.save(pptx_path)
print('Slide 8: Technology Stack compact grid generated successfully in PPTX!')
