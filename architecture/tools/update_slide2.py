import pptx
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

prs = pptx.Presentation(r'd:\olderEdrive\EshaProjects\AerifyAI\AerifyAI_AI_Builder_Cup_Pitch_Deck.pptx')
s2 = prs.slides[2]

# Remove previously added shapes beyond base 3 shapes (banner, bottom line, title)
while len(s2.shapes) > 3:
    sp = s2.shapes[-1]
    sp._element.getparent().remove(sp._element)

# Ensure title is clean and styled
for shape in s2.shapes:
    if shape.has_text_frame and 'Brief about the idea' in shape.text:
        tf = shape.text_frame
        p = tf.paragraphs[0]
        p.font.name = 'Google Sans'
        p.font.size = Pt(22)
        p.font.bold = True
        p.font.color.rgb = RGBColor(255, 255, 255)

# 4 Google-colored cards (2 columns x 2 rows)
# Exactly matching the colorful SIH reference style:
# Blue: #4285F4 | Yellow: #FBBC04 | Green: #34A853 | Red: #EA4335

col_lefts = [Emu(450000), Emu(4700000)]
row_tops = [Emu(1350000), Emu(3080000)]
card_w = Emu(4000000)
card_h = Emu(1580000)

google_cards = [
    {
        'col': 0, 'row': 0,
        'badge_icon': '🌐',
        'theme_color': RGBColor(66, 133, 244),  # Google Blue
        'card_bg': RGBColor(14, 26, 48),       # Dark Slate with Blue tint
        'title': 'Action-Driven Clean Air Platform',
        'subtitle': '01 // PLATFORM DEFINITION',
        'points': [
            'A Google-native municipal platform transforming passive pollution alerts into verified clean-air actions.',
            'Bridges satellite-assimilated 500m telemetry directly with on-the-ground municipal field operations.'
        ]
    },
    {
        'col': 1, 'row': 0,
        'badge_icon': '⚠️',
        'theme_color': RGBColor(251, 188, 4),  # Google Yellow
        'card_bg': RGBColor(38, 30, 14),       # Dark Slate with Yellow tint
        'title': 'Severe Smog & Unmeasured Spending',
        'subtitle': '02 // URBAN CONTEXT & NEED',
        'points': [
            'Megacity Smog Crisis: 30M+ citizens face hazardous seasonal pollution spikes every winter.',
            'Unverified Spend: Crores spent on anti-smog guns & sweepers with zero proof of actual impact.'
        ]
    },
    {
        'col': 0, 'row': 1,
        'badge_icon': '🎯',
        'theme_color': RGBColor(52, 168, 83),  # Google Green
        'card_bg': RGBColor(14, 36, 24),       # Dark Slate with Green tint
        'title': 'Verify Impact & Build Civic Trust',
        'subtitle': '03 // CORE MISSION & GOAL',
        'points': [
            'Scientific Attribution: Distinguish true clean-up results from ambient wind and weather changes.',
            'Civic Transparency: Show citizens verified clean-air actions in neighborhoods where they live & travel.'
        ]
    },
    {
        'col': 1, 'row': 1,
        'badge_icon': '👥',
        'theme_color': RGBColor(234, 67, 53),  # Google Red
        'card_bg': RGBColor(40, 18, 18),       # Dark Slate with Red tint
        'title': 'Citizens, Municipalities & Regulators',
        'subtitle': '04 // TARGET STAKEHOLDERS',
        'points': [
            'Municipal Bodies (MCD): Dispatch, schedule, and scientifically audit field interventions.',
            'Citizens & Commuters: Hyper-local street AQI, pollen risk, and 30-day verified action history.'
        ]
    }
]

for c in google_cards:
    left = col_lefts[c['col']]
    top = row_tops[c['row']]
    
    # Background Card with colorful border
    card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, card_w, card_h)
    card.fill.solid()
    card.fill.fore_color.rgb = c['card_bg']
    card.line.color.rgb = c['theme_color']
    card.line.width = Pt(1.75)
    
    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = Pt(14)
    tf.margin_top = Pt(10)
    tf.margin_right = Pt(14)
    tf.margin_bottom = Pt(8)
    
    # Subtitle / Category Tag with emoji badge
    p0 = tf.paragraphs[0]
    p0.text = f"{c['badge_icon']}  {c['subtitle']}"
    p0.font.name = 'Google Sans'
    p0.font.size = Pt(8.5)
    p0.font.bold = True
    p0.font.color.rgb = c['theme_color']
    p0.space_after = Pt(2)
    
    # Main Headline
    p1 = tf.add_paragraph()
    p1.text = c['title']
    p1.font.name = 'Google Sans'
    p1.font.size = Pt(12)
    p1.font.bold = True
    p1.font.color.rgb = RGBColor(255, 255, 255)
    p1.space_after = Pt(4)
    
    # Bullets
    for pt in c['points']:
        pb = tf.add_paragraph()
        pb.space_after = Pt(2)
        
        r_dot = pb.add_run()
        r_dot.text = '• '
        r_dot.font.name = 'Google Sans'
        r_dot.font.size = Pt(9.5)
        r_dot.font.bold = True
        r_dot.font.color.rgb = c['theme_color']
        
        r_txt = pb.add_run()
        r_txt.text = pt
        r_txt.font.name = 'Google Sans'
        r_txt.font.size = Pt(9.5)
        r_txt.font.color.rgb = RGBColor(228, 234, 245)

# Bottom Roadmap Banner (like in reference image)
b_left = Emu(450000)
b_top = Emu(4760000)
b_w = Emu(8250000)
b_h = Emu(270000)

banner = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, b_left, b_top, b_w, b_h)
banner.fill.solid()
banner.fill.fore_color.rgb = RGBColor(18, 25, 42)
banner.line.color.rgb = RGBColor(66, 133, 244)
banner.line.width = Pt(1)

btf = banner.text_frame
btf.margin_left = Pt(12)
btf.margin_top = Pt(4)
btf.margin_right = Pt(12)
btf.margin_bottom = Pt(4)

bp = btf.paragraphs[0]
r_tag = bp.add_run()
r_tag.text = '💡 Mission Scope | '
r_tag.font.name = 'Google Sans'
r_tag.font.size = Pt(9)
r_tag.font.bold = True
r_tag.font.color.rgb = RGBColor(66, 133, 244)

r_desc = bp.add_run()
r_desc.text = 'Active in Delhi-NCR (MCD) • Expanding to Mumbai (BMC) & National Industrial Corridors via Google Maps Environmental APIs'
r_desc.font.name = 'Google Sans'
r_desc.font.size = Pt(9)
r_desc.font.color.rgb = RGBColor(200, 215, 235)

prs.save(r'd:\olderEdrive\EshaProjects\AerifyAI\AerifyAI_AI_Builder_Cup_Pitch_Deck.pptx')
print('Slide 2 successfully formatted with colorful Google cards!')
